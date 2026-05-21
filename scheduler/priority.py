"""Priority URL ingestion - parse input en plan in voor push-runs.

Een priority URL is een URL die de gebruiker handmatig versneld geïndexeerd wil
zien. Priority URLs verbruiken het reguliere dagelijkse push-quotum (200/dag)
maar gaan voor in de wachtrij. Wanneer er meer dan 200 worden aangeboden voor
dezelfde dag worden ze automatisch over opeenvolgende dagen verdeeld.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from datetime import date
from pathlib import Path

from config.settings import INDEXING_DAILY_LIMIT
from config.shops import get_shop_for_url
from db.queries import insert_priority_urls

logger = logging.getLogger(__name__)

URL_PATTERN = re.compile(r"https?://[^\s,;\"']+", re.IGNORECASE)


def parse_text(content: str) -> list[str]:
    """Pak alle http(s) URLs uit een tekstblob (één per regel of vrij gemixt)."""
    urls: list[str] = []
    seen: set[str] = set()
    for match in URL_PATTERN.findall(content or ""):
        url = match.strip().rstrip(".,;:")
        if url and url not in seen:
            seen.add(url)
            urls.append(url)
    return urls


def parse_csv(content: str) -> list[str]:
    """Pak URLs uit CSV-content (elke cel kan een URL zijn)."""
    urls: list[str] = []
    seen: set[str] = set()
    reader = csv.reader(io.StringIO(content))
    for row in reader:
        for cell in row:
            if not cell:
                continue
            for match in URL_PATTERN.findall(cell):
                url = match.strip().rstrip(".,;:")
                if url and url not in seen:
                    seen.add(url)
                    urls.append(url)
    return urls


def parse_xlsx(file_path: str | Path) -> list[str]:
    """Pak URLs uit een .xlsx bestand. Werkt over alle sheets en cellen."""
    try:
        import openpyxl
    except ImportError as e:
        raise RuntimeError(
            "openpyxl is niet geïnstalleerd. Run: pip install openpyxl"
        ) from e

    wb = openpyxl.load_workbook(str(file_path), data_only=True, read_only=True)
    urls: list[str] = []
    seen: set[str] = set()
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if not cell:
                    continue
                cell_str = str(cell)
                for match in URL_PATTERN.findall(cell_str):
                    url = match.strip().rstrip(".,;:")
                    if url and url not in seen:
                        seen.add(url)
                        urls.append(url)
    return urls


def parse_file(file_path: str | Path) -> list[str]:
    """Detecteer bestandstype op extensie en parse."""
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return parse_xlsx(path)
    content = path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".csv":
        return parse_csv(content)
    return parse_text(content)


def schedule_priority_urls(
    urls: list[str],
    source: str = "manual",
    daily_limit: int | None = None,
    start_date: date | None = None,
) -> dict:
    """Plan een lijst URLs in als priority push-werk.

    URLs worden ge-deduped, gemapped op shop_id via het domein, en aan
    `insert_priority_urls` doorgegeven. Het effectieve daily_limit is standaard
    gelijk aan INDEXING_DAILY_LIMIT (200) — dat is meteen Google's hard cap.
    """
    if daily_limit is None:
        daily_limit = INDEXING_DAILY_LIMIT

    rows: list[tuple[str, str | None]] = []
    seen: set[str] = set()
    unmapped: list[str] = []
    for url in urls:
        if not url or url in seen:
            continue
        seen.add(url)
        shop = get_shop_for_url(url)
        if shop is None:
            unmapped.append(url)
        rows.append((url, shop.shop_id if shop else None))

    result = insert_priority_urls(
        rows,
        source=source,
        daily_limit=daily_limit,
        start_date=start_date,
    )
    result["received"] = len(urls)
    result["deduped"] = len(rows)
    result["unmapped"] = len(unmapped)
    result["unmapped_sample"] = unmapped[:5]
    return result
